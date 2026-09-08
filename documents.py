'''
    ******************************************************************************************
      Assembly:                Mappy
      Filename:                documents.py
      Author:                  Terry D. Eppler
      Created:                 09-08-2026
      Last Modified By:        Terry D. Eppler
      Last Modified On:        09-08-2026
    ******************************************************************************************
    <summary>
        Local document loaders used by Mappy document-processing workflows.
    </summary>
    ******************************************************************************************
'''
from __future__ import annotations
from io import BytesIO
from typing import List
import json

from bs4 import BeautifulSoup
from docx import Document as WordDocument
from langchain_core.documents import Document
import pandas as pd
from pypdf import PdfReader
from pptx import Presentation


def throw_if( name: str, value: object ) -> None:
	"""Validate a required runtime value.

	Purpose:
		Prevents local document-loading work from continuing with missing filenames, loader names,
		or file content.

	Args:
		name (str): Argument name included in validation errors.
		value (object): Runtime value to validate.

	Returns:
		None: This function validates input and does not return a value.
	"""
	if value is None:
		raise ValueError( f'Argument "{name}" cannot be empty!' )

	if isinstance( value, str ) and not value.strip( ):
		raise ValueError( f'Argument "{name}" cannot be empty!' )

	if isinstance( value, ( bytes, bytearray, list, tuple, dict, set ) ) and len( value ) == 0:
		raise ValueError( f'Argument "{name}" cannot be empty!' )


class DocumentLoaderFactory( ):
	"""Load supported local files into LangChain Document objects."""

	filename: str
	content: bytes
	loader_name: str

	def __init__( self ) -> None:
		"""Initialize loader state.

		Purpose:
			Initializes reusable document-loader state.

		Args:
			None.

		Returns:
			None: This method initializes instance state.
		"""
		self.filename = ''
		self.content = b''
		self.loader_name = ''

	def load( self, filename: str, content: bytes, loader_name: str ) -> List[ Document ]:
		"""Load file content using the selected local loader.

		Purpose:
			Converts Text, CSV, PDF, Excel, Word, Markdown, HTML, JSON, or PowerPoint content into
			LangChain Document objects with source metadata.

		Args:
			filename (str): Original uploaded filename.
			content (bytes): Uploaded file content.
			loader_name (str): Selected loader name.

		Returns:
			List[Document]: Loaded LangChain documents.
		"""
		throw_if( 'filename', filename )
		throw_if( 'content', content )
		throw_if( 'loader_name', loader_name )
		self.filename = filename
		self.content = content
		self.loader_name = loader_name

		if self.loader_name == 'Text Loader':
			return self.load_text( )
		if self.loader_name == 'CSV Loader':
			return self.load_csv( )
		if self.loader_name == 'PDF Loader':
			return self.load_pdf( )
		if self.loader_name == 'Excel Loader':
			return self.load_excel( )
		if self.loader_name == 'Word Document Loader':
			return self.load_word( )
		if self.loader_name == 'Markdown Loader':
			return self.load_markdown( )
		if self.loader_name == 'HTML Loader':
			return self.load_html( )
		if self.loader_name == 'JSON Loader':
			return self.load_json( )
		if self.loader_name == 'PowerPoint Loader':
			return self.load_powerpoint( )

		raise ValueError( f'Unsupported document loader: {self.loader_name}' )

	def load_text( self ) -> List[ Document ]:
		"""Load UTF-8 text content.

		Purpose:
			Converts uploaded text bytes into one LangChain document.

		Args:
			None.

		Returns:
			List[Document]: Loaded text document.
		"""
		text = self.content.decode( 'utf-8', errors='replace' )
		return [ Document( page_content=text, metadata={ 'source': self.filename } ) ]

	def load_csv( self ) -> List[ Document ]:
		"""Load CSV content.

		Purpose:
			Converts each CSV row to a LangChain document while preserving the row number.

		Args:
			None.

		Returns:
			List[Document]: One document for each CSV row.
		"""
		df_source = pd.read_csv( BytesIO( self.content ) )
		documents: List[ Document ] = [ ]
		for index, row in df_source.iterrows( ):
			text = '\n'.join( f'{column}: {row[ column ]}' for column in df_source.columns )
			documents.append( Document(
				page_content=text,
				metadata={ 'source': self.filename, 'row': int( index ) + 1 },
			) )
		return documents

	def load_pdf( self ) -> List[ Document ]:
		"""Load PDF pages.

		Purpose:
			Extracts text from each PDF page and creates one LangChain document per non-empty page.

		Args:
			None.

		Returns:
			List[Document]: Extracted PDF page documents.
		"""
		reader = PdfReader( BytesIO( self.content ) )
		documents: List[ Document ] = [ ]
		for index, page in enumerate( reader.pages, start=1 ):
			text = ( page.extract_text( ) or '' ).strip( )
			if text:
				documents.append( Document(
					page_content=text,
					metadata={ 'source': self.filename, 'page': index },
				) )
		return documents

	def load_excel( self ) -> List[ Document ]:
		"""Load Excel worksheets.

		Purpose:
			Converts each populated worksheet row to a LangChain document with sheet and row metadata.

		Args:
			None.

		Returns:
			List[Document]: One document for each populated Excel row.
		"""
		excel_file = pd.ExcelFile( BytesIO( self.content ) )
		documents: List[ Document ] = [ ]
		for sheet_name in excel_file.sheet_names:
			df_sheet = pd.read_excel( excel_file, sheet_name=sheet_name )
			for index, row in df_sheet.iterrows( ):
				text = '\n'.join( f'{column}: {row[ column ]}' for column in df_sheet.columns )
				documents.append( Document(
					page_content=text,
					metadata={
						'source': self.filename,
						'sheet': sheet_name,
						'row': int( index ) + 1,
					},
				) )
		return documents

	def load_word( self ) -> List[ Document ]:
		"""Load Word paragraphs.

		Purpose:
			Extracts non-empty paragraphs from a DOCX file into LangChain documents.

		Args:
			None.

		Returns:
			List[Document]: One document per non-empty paragraph.
		"""
		word = WordDocument( BytesIO( self.content ) )
		documents: List[ Document ] = [ ]
		for index, paragraph in enumerate( word.paragraphs, start=1 ):
			text = paragraph.text.strip( )
			if text:
				documents.append( Document(
					page_content=text,
					metadata={ 'source': self.filename, 'paragraph': index },
				) )
		return documents

	def load_markdown( self ) -> List[ Document ]:
		"""Load Markdown content.

		Purpose:
			Preserves uploaded Markdown text as one LangChain document for downstream chunking.

		Args:
			None.

		Returns:
			List[Document]: Loaded Markdown document.
		"""
		text = self.content.decode( 'utf-8', errors='replace' )
		return [ Document( page_content=text, metadata={ 'source': self.filename } ) ]

	def load_html( self ) -> List[ Document ]:
		"""Load visible HTML text.

		Purpose:
			Extracts visible textual content from uploaded HTML into one LangChain document.

		Args:
			None.

		Returns:
			List[Document]: Extracted HTML text document.
		"""
		html = self.content.decode( 'utf-8', errors='replace' )
		soup = BeautifulSoup( html, 'html.parser' )
		text = soup.get_text( '\n', strip=True )
		return [ Document( page_content=text, metadata={ 'source': self.filename } ) ]

	def load_json( self ) -> List[ Document ]:
		"""Load JSON content.

		Purpose:
			Parses uploaded JSON and preserves its structured representation as formatted text.

		Args:
			None.

		Returns:
			List[Document]: Loaded JSON document.
		"""
		data = json.loads( self.content.decode( 'utf-8', errors='replace' ) )
		text = json.dumps( data, indent=2, ensure_ascii=False )
		return [ Document( page_content=text, metadata={ 'source': self.filename } ) ]

	def load_powerpoint( self ) -> List[ Document ]:
		"""Load PowerPoint slide text.

		Purpose:
			Extracts text-bearing shapes from each PowerPoint slide into LangChain documents.

		Args:
			None.

		Returns:
			List[Document]: One document per slide containing text.
		"""
		presentation = Presentation( BytesIO( self.content ) )
		documents: List[ Document ] = [ ]
		for index, slide in enumerate( presentation.slides, start=1 ):
			parts: List[ str ] = [ ]
			for shape in slide.shapes:
				if hasattr( shape, 'text' ) and shape.text.strip( ):
					parts.append( shape.text.strip( ) )
			text = '\n'.join( parts ).strip( )
			if text:
				documents.append( Document(
					page_content=text,
					metadata={ 'source': self.filename, 'slide': index },
				) )
		return documents
